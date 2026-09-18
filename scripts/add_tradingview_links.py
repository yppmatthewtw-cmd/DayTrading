# -*- coding: utf-8 -*-
"""
為 output/ 內所有 Excel 的 ticker 加上 TradingView 超連結。

設計：
  1) 「短且只含一個 ticker」的儲存格 → 直接在原格加超連結（保留原有字體/底色，只轉藍色加底線）。
     長度上限避免把整段敘述文字變成連結。
  2) 含多個 ticker 的儲存格 → Excel 不支援在單一儲存格內為部分文字加連結，故不在原格處理，
     改由每個檔案新增的「TV_連結」索引表覆蓋（列出該檔案出現過的所有 ticker）。
  3) 冪等：可重複執行。已有連結的儲存格會略過；索引表每次重建。

用法:
    python3 scripts/add_tradingview_links.py                  # 處理 output/*.xlsx
    python3 scripts/add_tradingview_links.py <file> [...]     # 處理指定檔案
"""
import glob
import os
import re
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# 只有短標籤格才在原地加連結；再長的視為敘述文字，不動
MAX_INLINE_LEN = 30
INDEX_SHEET = "TV_連結"
LINK_COLOR = "0563C1"

# ticker → TradingView 交易所前綴
EXCHANGE = {
    "NASDAQ": """NVDA TSLA AAPL MSFT AMZN META GOOGL AVGO AMD MU INTC PLTR QCOM
                 NFLX CRWV IREN SOFI AMGN SMH SOXX""".split(),
    "NYSE": "ORCL PFE BE XOM CVX OXY SLB VLO DAL CRM HD GNRC CAT CMI".split(),
    "AMEX": "XLE XLU XLV XLF XLK SPY".split(),   # NYSE Arca ETF 在 TradingView 用 AMEX
    "CBOE": ["VIX"],
}
TICKERS = {t: f"{ex}:{t}" for ex, ts in EXCHANGE.items() for t in ts}

# 非 ticker 的指標名稱（儲存格文字須完全相符才連結）
NAME_MAP = {
    "Brent 原油": "TVC:UKOIL",
    "Brent 按年變動": "TVC:UKOIL",
    "Brent 期貨曲線形態": "TVC:UKOIL",
    "10 年期美債孳息": "TVC:US10Y",
    "標普 500": "SP:SPX",
    "S&P 500": "SP:SPX",
    "道指": "DJ:DJI",
    "納指": "NASDAQ:IXIC",
    "羅素2000": "TVC:RUT",
    "黃金": "TVC:GOLD",
}

NAMES = {
    "NVDA": "NVIDIA", "TSLA": "Tesla", "AAPL": "Apple", "MSFT": "Microsoft",
    "AMZN": "Amazon", "META": "Meta Platforms", "GOOGL": "Alphabet",
    "AVGO": "Broadcom", "AMD": "AMD", "MU": "Micron", "INTC": "Intel",
    "PLTR": "Palantir", "QCOM": "Qualcomm", "NFLX": "Netflix",
    "CRWV": "CoreWeave", "IREN": "IREN", "SOFI": "SoFi Technologies",
    "AMGN": "Amgen", "ORCL": "Oracle", "PFE": "Pfizer", "BE": "Bloom Energy",
    "XOM": "ExxonMobil", "CVX": "Chevron", "OXY": "Occidental Petroleum",
    "SLB": "SLB (Schlumberger)", "VLO": "Valero Energy", "DAL": "Delta Air Lines",
    "CRM": "Salesforce", "HD": "Home Depot",
    "SMH": "VanEck 半導體 ETF", "SOXX": "iShares 半導體 ETF",
    "XLE": "能源類股 ETF", "XLU": "公用事業 ETF", "XLV": "醫療保健 ETF",
    "XLF": "金融類股 ETF", "XLK": "科技類股 ETF", "SPY": "標普 500 ETF",
    "VIX": "CBOE 波動率指數",
    "TVC:UKOIL": "Brent 原油", "TVC:USOIL": "WTI 原油",
    "TVC:US10Y": "美國 10 年期孳息", "SP:SPX": "標普 500 指數",
    "DJ:DJI": "道瓊工業指數", "NASDAQ:IXIC": "納斯達克綜合指數",
    "TVC:RUT": "羅素 2000 指數", "TVC:GOLD": "黃金現貨",
}

TOKEN = re.compile(r"\b[A-Z]{1,5}\b")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def url_for(symbol):
    """TradingView 圖表連結（day trading 用圖表頁比概覽頁實用）。"""
    return "https://www.tradingview.com/chart/?symbol=" + symbol.replace(":", "%3A")


def link_font(cell):
    """保留原字體屬性，只改為連結外觀。"""
    f = cell.font
    return Font(name=f.name, size=f.size, bold=f.bold, italic=f.italic,
                color=LINK_COLOR, underline="single")


def tickers_in(text):
    return sorted(set(TOKEN.findall(text)) & TICKERS.keys())


def process(path):
    wb = load_workbook(path)
    if INDEX_SHEET in wb.sheetnames:          # 冪等：重建索引表
        del wb[INDEX_SHEET]

    inline = 0
    skipped_multi = 0
    seen = {}                                  # symbol -> [(sheet, coord), ...]

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                text = cell.value.strip()

                found = tickers_in(text)
                symbols = [TICKERS[t] for t in found]

                # 指標名稱（須完全相符），避免誤配敘述文字
                if not symbols and text in NAME_MAP:
                    symbols = [NAME_MAP[text]]

                if not symbols:
                    continue

                for s in symbols:
                    seen.setdefault(s, []).append(f"{ws.title}!{cell.coordinate}")

                if len(symbols) > 1:
                    skipped_multi += 1        # 單格內無法為部分文字加連結
                    continue
                if len(text) > MAX_INLINE_LEN or cell.hyperlink is not None:
                    continue

                cell.hyperlink = url_for(symbols[0])
                cell.font = link_font(cell)
                inline += 1

    build_index(wb, seen, inline, skipped_multi)
    wb.save(path)
    return inline, skipped_multi, len(seen)


def build_index(wb, seen, inline, skipped_multi):
    ws = wb.create_sheet(INDEX_SHEET)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 46

    t = ws.cell(row=1, column=1, value="TradingView 連結索引（本檔案出現過的所有代號）")
    t.font = Font(bold=True, size=14, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor="1F3864")
    t.alignment = Alignment(vertical="center")
    ws.merge_cells("A1:D1")
    ws.row_dimensions[1].height = 24

    n = ws.cell(row=2, column=1, value=(
        f"原地已加連結的儲存格：{inline} 個。"
        f"另有 {skipped_multi} 個儲存格同時提及多個代號 —— Excel 不支援在單一儲存格內為部分文字加連結，"
        "故該類儲存格不在原地處理，改由本表覆蓋。點擊下方連結即可開啟該代號的 TradingView 圖表。"))
    n.font = Font(italic=True, size=9, color="595959")
    n.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:D2")
    ws.row_dimensions[2].height = 32

    for i, h in enumerate(["代號", "名稱", "TradingView", "出現位置"], start=1):
        c = ws.cell(row=4, column=i, value=h)
        c.font = Font(bold=True, size=10, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F5597")
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 20

    r = 5
    for sym in sorted(seen, key=lambda s: (":" in s and s.split(":")[0] not in
                                           ("NASDAQ", "NYSE", "AMEX", "CBOE"), s.split(":")[-1])):
        short = sym.split(":")[-1]
        locs = seen[sym]
        where = "、".join(locs[:6]) + (f" …等 {len(locs)} 處" if len(locs) > 6 else "")

        ws.cell(row=r, column=1, value=short).font = Font(size=10, bold=True)
        ws.cell(row=r, column=2, value=NAMES.get(short) or NAMES.get(sym, ""))
        c = ws.cell(row=r, column=3, value=sym)
        c.hyperlink = url_for(sym)
        c.font = Font(size=10, color=LINK_COLOR, underline="single")
        ws.cell(row=r, column=4, value=where).font = Font(size=9, color="595959")

        for cc in range(1, 5):
            ws.cell(row=r, column=cc).border = BOX
            ws.cell(row=r, column=cc).alignment = Alignment(vertical="top", wrap_text=(cc == 4))
            if (r - 5) % 2:
                ws.cell(row=r, column=cc).fill = PatternFill("solid", fgColor="F2F5FB")
        ws.row_dimensions[r].height = 18
        r += 1


def main():
    paths = sys.argv[1:] or sorted(glob.glob(
        os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                     "output", "*.xlsx")))
    if not paths:
        sys.exit("找不到任何 .xlsx 檔案。")
    for p in paths:
        inline, multi, uniq = process(p)
        print(f"{os.path.basename(p)}\n"
              f"    原地連結 {inline} 格 | 多代號儲存格 {multi} 格（由索引表覆蓋）| 不重複代號 {uniq} 個")


if __name__ == "__main__":
    main()
