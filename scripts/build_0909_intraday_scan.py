# -*- coding: utf-8 -*-
"""2026-09-09 盤中掃描 (美東 ~11:12, 開市約 1h42m)：低波幅 + 20 分鐘均線平穩向上"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference

OUT=os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),"output")
os.makedirs(OUT,exist_ok=True)
STAMP="2026-09-09 11:12 ET (港時 23:12)"

NAVY="1F3864"; SLATE="2F5597"
RED="FFC7CE"; AMBER="FFEB9C"; GREEN="C6EFCE"
REDF="9C0006"; AMBERF="9C6500"; GREENF="006100"
THIN=Side(style="thin",color="BFBFBF"); BOX=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)

def title_row(ws,r,t,span,size=14):
    c=ws.cell(row=r,column=1,value=t); c.font=Font(bold=True,size=size,color="FFFFFF")
    c.fill=PatternFill("solid",fgColor=NAVY); c.alignment=Alignment(vertical="center",wrap_text=True)
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=span); ws.row_dimensions[r].height=26
def note_row(ws,r,t,span,h=34):
    c=ws.cell(row=r,column=1,value=t); c.font=Font(italic=True,size=9,color="595959")
    c.alignment=Alignment(wrap_text=True,vertical="top")
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=span); ws.row_dimensions[r].height=h
def header(ws,r,hs,h=32):
    for i,x in enumerate(hs,1):
        c=ws.cell(row=r,column=i,value=x); c.font=Font(bold=True,size=10,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor=SLATE); c.border=BOX
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    ws.row_dimensions[r].height=h
def widths(ws,spec):
    for k,v in spec.items(): ws.column_dimensions[k].width=v
def put(ws,r,row,wrapcols=(),height=26,fill=None):
    for i,v in enumerate(row,1):
        c=ws.cell(row=r,column=i,value=v); c.border=BOX; c.font=Font(size=10)
        c.alignment=Alignment(vertical="top",wrap_text=(i in wrapcols))
        if fill: c.fill=PatternFill("solid",fgColor=fill)
    ws.row_dimensions[r].height=height

wb=Workbook()

# ============ 00 掃描結論 ============
ws=wb.active; ws.title="00_掃描結論"
widths(ws,{"A":6,"B":9,"C":11,"D":11,"E":11,"F":56})
title_row(ws,1,f"盤中掃描：波幅小 + 20 分鐘均線平穩向上　|　{STAMP}",6)
note_row(ws,2,("★ 今日與昨日 (9/8) 的結論有一個重要改變：AAPL 昨日是「低波幅基準」，"
 "今日因為有產品發布會（Surprise and Shine，iPhone 18 Pro + 首款摺疊機，新任 CEO Ternus 首場）"
 "而變成事件股 —— 盤中的排程事件正是打斷 20 分鐘均線平穩度的典型原因，今日必須剔除。"),6,44)

title_row(ws,4,"今日大市與板塊（已核實）",6,12)
put(ws,5,["指數","S&P 500","道指","納指","羅素2000","板塊"],height=20,fill="EDF1F9")
put(ws,6,["變動","-0.24%","-0.58%","-0.35%","-0.52%",
 "領先：SMH +1.19%（半導體）、XLE +1.11%（能源）、XLU +0.86%（公用）｜落後：XLV -2.52%、XLF -1.38%"],
 wrapcols=(6,),height=30)

title_row(ws,8,"篩選結果（由最符合排起）",6,12)
header(ws,9,["排名","代號","今日變動","波幅級別","評分","理由 / 注意事項"])
sel=[
 [1,"XOM","+2.30%","低（約1.5-2%）","★★★★★",
  "★ 今日最佳。三個條件同時成立：\n"
  "① 綜合型油氣巨企，油價 beta 遠低於 E&P 與油服 → 日內振幅是能源板塊中最小的；\n"
  "② 驅動力是原油「整個時段逐步推升穿破 $100」，屬持續性買盤而非單點跳空 → 20 分鐘均線最可能呈平穩上行；\n"
  "③ XLE +1.11% 印證是板塊性資金流，非個股炒作。今日無自身事件風險。"],
 [2,"CVX","+2.10%","低（約1.8-2.2%）","★★★★☆",
  "與 XOM 同一邏輯，波幅略高一線。可作為 XOM 的替代或分散。"],
 [3,"AVGO","隨 SMH 走強","中（約2.5-3.5%）","★★★★☆",
  "SMH +1.19% 領先全場，延續昨日的板塊性穩步走高型態（昨日 +2.98%，無消息跳空）。\n"
  "注意：本週有高盛科技會議，NVDA / AVGO 均在名單上 → 隨時可能有會議發言引發的跳動。"],
 [4,"NVDA","隨 SMH 走強","低（約1.8-2.2%）","★★★☆☆",
  "全市場最厚的交投 → 相對波幅最小、跳動最平滑，是最乾淨的 20 分鐘均線載體。\n"
  "注意：同樣有高盛科技會議的頭條風險，斜率通常較 AVGO 平緩。"],
 [5,"XLU 成分","+0.86%（板塊）","最低","★★★☆☆",
  "公用事業是全市場波幅最低的板塊，今日逆市收升 → 均線最平滑。\n"
  "注意：① 純公用股一般擠不進成交額 20 大；② 在 10 年期創 2023 年底以來新高的環境下公用股上升，"
  "屬防守性輪動，若 13:00 標售順利、債息回落，這個資金流可能瞬間反轉。"],
]
r=10
for row in sel:
    put(ws,r,row,wrapcols=(6,),height=max(30,16*row[5].count("\n")+34),fill="EAF3EA")
    ws.cell(row=r,column=5).font=Font(size=11,bold=True,color=GREENF)
    ws.cell(row=r,column=5).alignment=Alignment(horizontal="center",vertical="center")
    ws.cell(row=r,column=4).fill=PatternFill("solid",fgColor=GREEN)
    r+=1

r+=1
title_row(ws,r,"明確剔除（今日）",6,12); r+=1
header(ws,r,["代號","今日情況","","","","剔除理由"],h=20); r+=1
exc=[
 ["AAPL","產品發布會（盤中）","★ 昨日是低波幅首選，今日剔除。排程事件會在盤中製造跳動，正正打斷 20 分鐘均線的平穩度。市場共識亦是「fade 消費科技的反彈」"],
 ["META","Muse AI agent 發布","消息跳空後急升 —— 是 gap，不是平穩上行。屬另一種交易（動能），不符合本次篩選條件"],
 ["VLO / 煉油股","裂解價差擴闊","分析師的最高信心多倉，但裂解價差股是油價的高 beta 放大器，波幅遠大於綜合型"],
 ["OXY / SLB / E&P","油價高 beta","同上，日內振幅通常是 XOM 的 2 倍"],
 ["DAL / 運輸股","燃油成本受壓","方向向下"],
 ["XLV 成分","板塊 -2.52%","全場最弱，方向不符"],
 ["XLF 成分","板塊 -1.38%","次弱，方向不符"],
 ["TSLA / MU / PLTR / CRWV / IREN","—","結構性高波幅，任何時候都不符合「波幅小」"],
]
for row in exc:
    put(ws,r,[row[0],row[1],"","","",row[2]],wrapcols=(6,),height=24,fill="FDECEC")
    ws.cell(row=r,column=1).font=Font(size=10,bold=True)
    r+=1

r+=1
title_row(ws,r,"★ 未來 3 小時最大的單一風險：13:00 ET 的 $39B 十年期國債標售",6,12); r+=1
c=ws.cell(row=r,column=1,value=(
 "現在是 11:12 ET，距離標售約 1 小時 48 分鐘。10 年期孳息已升至 2023 年底以來最高。\n\n"
 "若標售結果疲弱（尾部利差擴大、投標倍數偏低）→ 債息急升 → 全市場的 20 分鐘均線會在同一分鐘內一起轉向，"
 "包括上面所有「平穩向上」的候選。這不是個股風險，是系統性風險，分散持股完全無法對沖。\n\n"
 "操作含意：無論選中哪一隻，13:00 之前應已了結或至少減至 1/3 倉位。"
 "標售結果通常在 13:01–13:03 之間反映在債價上，股市在 13:05 前跟隨。"))
c.font=Font(size=10.5); c.alignment=Alignment(vertical="top",wrap_text=True); c.border=BOX
c.fill=PatternFill("solid",fgColor=AMBER)
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6); ws.row_dimensions[r].height=104
r+=2

title_row(ws,r,"一句話總結",6,13); r+=1
c=ws.cell(row=r,column=1,value=(
 "今日的答案是 XOM（其次 CVX），不是昨日的 AAPL。\n\n"
 "原因是驅動力的「形狀」變了：昨日 AAPL 靠的是低 beta 的靜止，今日它有事件；"
 "而今日原油整個時段逐步推穿 $100，是一種持續、漸進、無跳空的買盤 —— 這正是製造「平穩向上的 20 分鐘均線」"
 "的唯一一種資金流形態。XOM 是這股資金流當中波幅最低的載體。\n\n"
 "若只想留在科技股：AVGO（斜率較清晰）或 NVDA（波幅較小），但要接受高盛科技會議的頭條風險。"))
c.font=Font(size=11); c.alignment=Alignment(vertical="top",wrap_text=True); c.border=BOX
ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=6); ws.row_dimensions[r].height=110

# ============ 01 板塊圖 ============
ws=wb.create_sheet("01_板塊資金流")
widths(ws,{"A":18,"B":12,"C":10,"D":52})
title_row(ws,1,f"板塊資金流（已核實）　|　{STAMP}",4)
header(ws,3,["板塊 ETF","今日變動","方向","解讀"])
sec=[["SMH 半導體",0.0119,"領先","AI 資本開支敘事未斷；NVDA/AVGO 在內。板塊性買盤＝均線平穩度最佳的一類"],
 ["XLE 能源",0.0111,"領先","Brent 兩個月來首次穿破 $100。持續性、漸進式買盤，最適合順勢做多低波幅標的"],
 ["XLU 公用",0.0086,"領先","波幅最低的板塊，逆市上升＝防守性輪動。但與債息上升同時發生，資金流不穩定"],
 ["XLV 醫療",-0.0252,"最弱","全場最弱。昨日 Amgen -10.15% 的餘波 + 板塊輪出"],
 ["XLF 金融",-0.0138,"弱","次弱。債息上升本應利好，但衰退/信貸成本疑慮壓過"]]
r=4
for row in sec:
    put(ws,r,row,wrapcols=(4,),height=28)
    ws.cell(row=r,column=2).number_format="+0.00%;-0.00%"
    f=GREEN if row[1]>0 else RED
    ws.cell(row=r,column=3).fill=PatternFill("solid",fgColor=f)
    ws.cell(row=r,column=3).font=Font(size=10,bold=True,color=GREENF if row[1]>0 else REDF)
    ws.cell(row=r,column=3).alignment=Alignment(horizontal="center",vertical="center")
    r+=1
ch=BarChart(); ch.type="bar"; ch.style=11
ch.title="2026-09-09 盤中板塊表現：只有 3 個板塊是綠的"
ch.height=8; ch.width=22; ch.legend=None
ch.add_data(Reference(ws,min_col=2,min_row=3,max_row=r-1),titles_from_data=True)
ch.set_categories(Reference(ws,min_col=1,min_row=4,max_row=r-1))
ws.add_chart(ch,f"A{r+2}")

# ============ 02 今日背景 ============
ws=wb.create_sheet("02_今日背景與時間表")
widths(ws,{"A":16,"B":72})
title_row(ws,1,"今日的宏觀背景（為何是這幾隻）",2)
bg=[
 ["原油破 $100","Brent 兩個月來首次穿破 $100；WTI 24 小時內 +6% 至約 $91.94。"
  "起因：美軍在 Kharg 島附近擊沉/擊中 5 艘伊朗油輪，此前伊朗向美軍設施發射彈道導彈。\n"
  "★ 這正好觸發了本 repo 四力框架中「Brent > $100 = 通脹二次上行確認」的紅色閾值。"],
 ["加息預期","交易員已將 9 月聯儲局的預期由「觀望」轉向「加息」。上週五（9/4）的 8 月非農偏強，加深了這個轉向。"],
 ["債息","10 年期孳息升至 2023 年底以來最高。今日 13:00 ET 標售 390 億美元 10 年期票據。"],
 ["標普位置","昨日收 7,673.52，今日 -0.24% 約 7,655。距離框架中的關鍵支撐 7,600 只有約 55 點 / 0.72%。\n"
  "★ 跌破 7,600 會啟動 CTA 與波動率控制基金的機械式減倉。"],
 ["個股事件","AAPL 產品發布會（iPhone 18 Pro + 首款摺疊機，新 CEO Ternus 首場）；"
  "META 發布 Muse 個人 AI agent，股價跳升；NVDA / AVGO 出席高盛科技會議。"],
 ["未來 2 日","9/11（五）8 月 CPI —— 這是驗證「油價滲入通脹」的關鍵數據，亦是本週最大的事件風險。"],
]
r=3
for a,b in bg:
    ws.cell(row=r,column=1,value=a).font=Font(size=10,bold=True); ws.cell(row=r,column=1).border=BOX
    ws.cell(row=r,column=1).alignment=Alignment(vertical="top",wrap_text=True)
    c=ws.cell(row=r,column=2,value=b); c.font=Font(size=10); c.border=BOX
    c.alignment=Alignment(vertical="top",wrap_text=True)
    ws.row_dimensions[r].height=max(34,17*b.count("\n")+36)
    r+=1

r+=1
title_row(ws,r,"今日剩餘時段的時間表（美東時間）",2); r+=1
header(ws,r,["時間","事件與應對"]); r+=1
for a,b in [
 ["11:12","★ 現在。開市 1 小時 42 分鐘。開盤區間已確立，午盤流動性即將轉薄"],
 ["11:30–13:00","午盤：流動性最低、假突破最多。若已持倉則收緊停損，不宜新開倉"],
 ["13:00","★ $39B 十年期標售。全日最大的單一風險。之前應已了結或減至 1/3"],
 ["13:05–14:00","標售消化。若債息急升 → 所有平穩上行的均線同時轉向；若順利 → 公用/防守股的資金流反轉"],
 ["約 13:00 起","AAPL 發布會（通常美西 10:00 開始）→ 消費科技頭條密集"],
 ["15:00–16:00","尾盤。昨日標普收於全日最低，今日若重演，15:30 後不應持有多倉"],
 ["盤後 / 隔夜","中東頭條的主要發生時段（亞洲盤）→ 原油與能源股的隔夜跳空風險最高"],
]:
    put(ws,r,[a,b],wrapcols=(2,),height=22)
    ws.cell(row=r,column=1).font=Font(size=10,bold=True)
    r+=1

# ============ 03 限制與來源 ============
ws=wb.create_sheet("03_限制與來源")
widths(ws,{"A":24,"B":76})
title_row(ws,1,"資料限制與來源",2)
items=[
 ["★ 資料限制","本執行環境的網絡政策封鎖了全部行情 API（Yahoo Finance、Stooq、Polygon、Alpha Vantage、"
  "TwelveData、FMP、marketdata.app、IEX、Nasdaq、CNBC 全部回 403），已於本次再度逐一測試確認。\n"
  "因此我無法取得任何逐分鐘 K 線，也就無法真正計算 20 分鐘移動平均線的斜率與 R²。\n"
  "本檔案中：指數變動、板塊 ETF 變動、XOM/CVX 的個股變動、各項事件 —— 均為已核實的公開資料；"
  "「波幅級別」與「20 分鐘均線是否平穩向上」為依板塊資金流形態、驅動力性質（持續推升 vs 消息跳空）"
  "與各股常態波動特性所作的推斷。請勿當作實測數據使用。"],
 ["如何取得實測結果","scripts/screen_20min_ma.py 實作了完整邏輯，在有行情權限的環境執行即可：\n"
  "  python3 scripts/screen_20min_ma.py --date 2026-09-09 --top 20 --out scan.csv\n"
  "判定門檻：rising_ratio ≥ 0.70、迴歸斜率 > 0、R² ≥ 0.60，通過者再以日內波幅由小至大排序。"],
 ["本次推斷的核心邏輯","「平穩向上的 20 分鐘均線」只由一種資金流形態產生：持續、漸進、無跳空的買盤。\n"
  "  · 板塊性資金流（XLE 隨油價逐步推升、SMH 隨 AI 敘事）→ 產生平穩上行 ✔\n"
  "  · 單一消息跳空（META 的 Muse、昨日 QCOM 的 AWS 合作）→ 產生階梯或沖高回落 ✘\n"
  "  · 排程事件（AAPL 今日發布會）→ 在事件時點打斷均線 ✘\n"
  "因此在同一股資金流之中，選「波幅最低的載體」即為答案 —— 能源流中是 XOM，半導體流中是 NVDA。"],
 ["免責聲明","本文件為市場資料整理與交易框架研究，非投資建議、非買賣要約。"
  "日內交易涉及高風險並可能損失全部本金。所有價位與判斷請以即時報價覆核。"],
]
r=3
for a,b in items:
    ws.cell(row=r,column=1,value=a).font=Font(size=10,bold=True); ws.cell(row=r,column=1).border=BOX
    ws.cell(row=r,column=1).alignment=Alignment(vertical="top",wrap_text=True)
    c=ws.cell(row=r,column=2,value=b); c.font=Font(size=10); c.border=BOX
    c.alignment=Alignment(vertical="top",wrap_text=True)
    ws.row_dimensions[r].height=max(50,16*b.count("\n")+52)
    r+=1

r+=1
header(ws,r,["項目","來源"]); r+=1
for a,b in [
 ["指數變動 (9/9)","TheStreet《Stock Market Today (Sept. 9, 2026)》；Yahoo Finance 9/9 live"],
 ["板塊 ETF (SMH/XLE/XLU/XLV/XLF)","Trading Strategy Guides《Stock Market Preview September 9, 2026》"],
 ["Brent 破 $100","Yahoo Finance《Dow, S&P 500, Nasdaq slip as oil prices hit $100》；TheStreet"],
 ["XOM +2.3% / CVX +2.1%","Benzinga《Energy Stocks Rally As Crude Spikes》"],
 ["Kharg 島油輪事件","TheStreet 9/9；Yahoo Finance 9/9"],
 ["AAPL 發布會","CNBC《What's expected at Apple's September event》；24/7 Wall St.；TradingKey"],
 ["META Muse","Trading Strategy Guides《5 Stocks To Watch On September 9, 2026》"],
 ["10Y 孳息 / $39B 標售","Yahoo Finance 9/9；Investrade《Mid-Morning Look: September 09, 2026》"],
 ["加息預期轉向","Yahoo Finance 9/9；上週五 8 月非農偏強"],
]:
    put(ws,r,[a,b],wrapcols=(1,2),height=22)
    ws.cell(row=r,column=1).font=Font(size=10,bold=True)
    r+=1

p=os.path.join(OUT,"2026-09-09_盤中掃描_低波幅20分鐘均線.xlsx")
wb.save(p); print("XLSX ->",p)
